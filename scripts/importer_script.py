import argparse
import hashlib
import json
import logging
import os
import re
import shutil
import ssl
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, time as dt_time, timedelta
from pathlib import Path

try:
    import xlrd
except ImportError:
    xlrd = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

REPO_ROOT = Path(__file__).resolve().parent.parent
LOGGER = logging.getLogger("ddo_importer")


def _load_env_file(path: Path) -> None:
    if not path.exists():
        return
    for raw in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


_load_env_file(REPO_ROOT / ".env")

CONFIG = {
    "location": os.getenv("LOCATION_CODE", "NOIDA").strip(),
    "api_endpoint": os.getenv("DDO_API_ENDPOINT", "").strip(),
    "api_token": (
        os.getenv("ATTENDANCE_INTEGRATION_TOKEN") or os.getenv("DDO_API_TOKEN") or ""
    ).strip(),
    "api_timeout_seconds": int(os.getenv("DDO_API_TIMEOUT_SECONDS", "30")),
    "api_max_retries": int(os.getenv("DDO_API_MAX_RETRIES", "3")),
    # Small chunks + pause between POSTs so DDO++ / Postgres can keep up.
    "api_chunk_size": int(os.getenv("DDO_API_CHUNK_SIZE", "50")),
    "api_chunk_delay_seconds": float(os.getenv("DDO_API_CHUNK_DELAY_SECONDS", "2")),
    "api_max_body_bytes": int(os.getenv("DDO_API_MAX_BODY_BYTES", "90000")),
    "attendance_dir": os.getenv("ATTENDANCE_DIR", r"D:\Attendance").strip(),
    "processed_dir": str(REPO_ROOT / "processed"),
    "failed_dir": str(REPO_ROOT / "failed"),
    "output_dir": str(REPO_ROOT / "output"),
    "output_batch_keep": int(os.getenv("OUTPUT_BATCH_KEEP", "10")),
    "archive_keep": int(os.getenv("ARCHIVE_KEEP", "15")),
    "stable_seconds": int(os.getenv("DROP_STABLE_SECONDS", "15")),
}

ALLOWED_LOCATIONS = {
    "AGRA": "AGRA",
    "NOIDA": "NOIDA",
    "HYD": "HYD",
    "HYDERABAD": "HYD",
}
LOCATION_FOLDERS = ("AGRA", "NOIDA", "HYD")
EXCEL_SUFFIXES = {".xls", ".xlsx"}

STATUS_MAP = {
    "P": "PRESENT",
    "A": "ABSENT",
    "WO": "WEEKLY_OFF",
    "CL": "CASUAL_LEAVE",
    "EL": "EARNED_LEAVE",
    "WFH": "WORK_FROM_HOME",
    "CO": "COMP_OFF",
    "LOP": "LOSS_OF_PAY",
    "CLP": "CASUAL_LEAVE",
    "ELP": "EARNED_LEAVE",
    "EFH": "WORK_FROM_HOME",
}

DEPARTMENT_NAMES = {
    "delivery", "ta", "sales & bd", "marketing", "hr", "f&a", "management", "it", "default",
}

WEEKDAY_CODES = {
    "M": 0,
    "T": 1,
    "W": 2,
    "TH": 3,
    "F": 4,
    "ST": 5,
    "S": 6,
}


@dataclass
class DayRecord:
    date: date
    status: str
    raw_status: str
    header: str
    in_time: str | None = None
    out_time: str | None = None
    total_minutes: int | None = None


@dataclass
class EmployeeBlock:
    code: str
    name: str
    days: list


@dataclass
class Employee:
    code: str
    name: str
    location: str


@dataclass
class AttendanceRecord:
    employee_code: str
    employee_name: str
    location: str
    date: str
    status: str
    raw_status: str
    in_time: str | None = None
    out_time: str | None = None
    total_minutes: int | None = None

    def to_api_dict(self) -> dict:
        return {
            "employee_code": self.employee_code,
            "employee_name": self.employee_name,
            "attendance_date": self.date,
            "status": self.raw_status,
            "in_time": self.in_time,
            "out_time": self.out_time,
            "total_minutes": self.total_minutes,
        }


@dataclass
class ImportBatch:
    id: str
    file_name: str
    location: str
    file_hash: str = ""
    status: str = "PENDING"
    error: str | None = None
    period: dict = field(default_factory=dict)
    failures: list = field(default_factory=list)
    data_quality_issues: list = field(default_factory=list)
    attendance_rows: list = field(default_factory=list)
    api_result: dict = field(default_factory=dict)


_current_batch: ImportBatch | None = None
_output_dir: Path | None = None
_employee_names: dict[str, str] = {}
_unknown_employee_names: list[str] = []


def _configure_logging() -> None:
    if LOGGER.handlers:
        return
    log_dir = Path(CONFIG["output_dir"])
    log_dir.mkdir(parents=True, exist_ok=True)
    formatter = logging.Formatter("%(asctime)s %(levelname)s %(message)s")
    LOGGER.setLevel(logging.INFO)
    stream = logging.StreamHandler(sys.stdout)
    stream.setFormatter(formatter)
    file_handler = logging.FileHandler(log_dir / "importer.log", encoding="utf-8")
    file_handler.setFormatter(formatter)
    LOGGER.addHandler(stream)
    LOGGER.addHandler(file_handler)
    LOGGER.propagate = False


def get_location_from_config():
    return os.getenv("LOCATION_CODE", CONFIG.get("location", "")).strip()


def normalize_location(location):
    code = (location or "").strip().upper()
    mapped = ALLOWED_LOCATIONS.get(code)
    if not mapped:
        raise ValueError("Location must be AGRA, NOIDA, or HYD")
    return mapped


def validate_location(location):
    if not location:
        raise ValueError("Location is not configured")
    return normalize_location(location)


def create_import_batch(file, location):
    global _current_batch, _output_dir, _employee_names, _unknown_employee_names
    _employee_names = {}
    _unknown_employee_names = []
    _output_dir = Path(CONFIG["output_dir"])
    _output_dir.mkdir(parents=True, exist_ok=True)
    Path(CONFIG["processed_dir"]).mkdir(parents=True, exist_ok=True)
    Path(CONFIG["failed_dir"]).mkdir(parents=True, exist_ok=True)

    batch = ImportBatch(
        id=str(uuid.uuid4())[:8],
        file_name=Path(file).name,
        location=location,
        file_hash=_file_hash(Path(file)),
    )
    _current_batch = batch
    LOGGER.info("Created import batch %s for %s (%s)", batch.id, batch.file_name, location)
    return batch


def read_excel(file):
    path = Path(file)
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        if load_workbook is None:
            raise RuntimeError("Missing dependency: openpyxl. Install with: python -m pip install openpyxl")
        sheets = _read_xlsx_sheets(path)
    elif suffix == ".xls":
        if xlrd is None:
            raise RuntimeError("Missing dependency: xlrd. Install with: python -m pip install xlrd")
        sheets = _read_xls_sheets(path)
    else:
        raise ValueError(f"Unsupported Excel type: {suffix}")

    sheet = _select_sheet(sheets)
    return {
        "file_name": path.name,
        "sheet_name": sheet["name"],
        "rows": sheet["rows"],
    }


def _read_xls_sheets(path):
    workbook = xlrd.open_workbook(path)
    sheets = []
    for sheet in workbook.sheets():
        rows = [
            [_cell_text(sheet.cell_value(row_idx, col_idx)) for col_idx in range(sheet.ncols)]
            for row_idx in range(sheet.nrows)
        ]
        sheets.append({"name": sheet.name, "rows": rows})
    return sheets


def _read_xlsx_sheets(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheets = []
        for sheet in workbook.worksheets:
            rows = [[_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            sheets.append({"name": sheet.title, "rows": rows})
        return sheets
    finally:
        workbook.close()


def extract_report_period(report):
    if report.get("_period"):
        return report["_period"]

    rows = report["rows"]
    header_row_index = next((idx for idx, row in enumerate(rows) if any(_parse_day_header(cell) for cell in row)), 0)
    headers = rows[header_row_index] if rows else []
    day_columns = [idx for idx, cell in enumerate(headers) if _parse_day_header(cell)]
    day_headers = [headers[idx] for idx in day_columns]
    report_text = " ".join(cell for row in rows[:8] for cell in row)
    year, month = _infer_year_month(report["file_name"], report_text)
    dates = _map_headers_to_dates(day_headers, year, month)
    first_day_col = day_columns[0] if day_columns else 0
    period = {
        "label": f"{datetime(year, month, 1):%B %Y}",
        "year": year,
        "month": month,
        "start": dates[0].isoformat() if dates else None,
        "end": dates[-1].isoformat() if dates else None,
        "dates": dates,
        "day_headers": day_headers,
        "day_columns": day_columns,
        "header_offset": max(first_day_col - 1, 0),
        "header_row_index": header_row_index,
        "has_code_header": any(re.search(r"emp\.?\s*code", str(cell), re.I) for cell in headers),
    }
    if _current_batch is not None:
        _current_batch.period = {
            "label": period["label"],
            "start": period["start"],
            "end": period["end"],
        }
    report["_period"] = period
    LOGGER.info("Report period: %s (%s to %s)", period["label"], period["start"], period["end"])
    return period


def extract_employee_blocks(report):
    rows = report["rows"]
    if not rows:
        return []

    period = extract_report_period(report)
    punch_format = any(_row_label(row) for row in rows)
    blocks = _extract_punch_blocks(rows, period) if punch_format else _extract_grid_blocks(rows, period)

    for block in blocks:
        if block.code:
            _employee_names[block.code] = block.name
        elif block.name:
            block.code = _unmapped_code(block.name)
            _employee_names[block.code] = block.name
            _unknown_employee_names.append(block.name)
            if _current_batch is not None:
                _current_batch.data_quality_issues.append(
                    {
                        "employee_code": block.code,
                        "employee_name": block.name,
                        "date": None,
                        "reason": "Missing employee_code; name-based identity used",
                    }
                )

    LOGGER.info("Found %s employee blocks in %s", len(blocks), report["sheet_name"])
    return blocks


def resolve_employee(code, location):
    if not code:
        return None
    employee_code = str(code).strip()
    return Employee(
        code=employee_code,
        name=_employee_names.get(employee_code, ""),
        location=location,
    )


def record_failure(batch, employee_code, reason):
    name = _employee_names.get(employee_code or "", "")
    if not employee_code and _unknown_employee_names:
        name = _unknown_employee_names.pop(0)
    item = {
        "employee_code": employee_code or "",
        "employee_name": name,
        "reason": reason,
    }
    batch.failures.append(item)
    LOGGER.warning("Failure: %s (%s)", reason, employee_code or name or "no code")


def record_data_quality_issue(batch, employee_code, day_date):
    item = {
        "employee_code": employee_code,
        "date": day_date.isoformat() if hasattr(day_date, "isoformat") else str(day_date),
        "reason": "Missing attendance status",
    }
    batch.data_quality_issues.append(item)
    LOGGER.warning("Data quality: missing status for %s on %s", employee_code, item["date"])


def transform(day, employee, location):
    return AttendanceRecord(
        employee_code=employee.code,
        employee_name=employee.name,
        location=location,
        date=day.date.isoformat(),
        status=STATUS_MAP.get(day.status, day.status),
        raw_status=day.raw_status,
        in_time=day.in_time,
        out_time=day.out_time,
        total_minutes=day.total_minutes,
    )


def upsert_attendance(attendance, batch_id):
    if _current_batch is None or _current_batch.id != batch_id:
        raise RuntimeError(f"No active batch {batch_id}")
    _current_batch.attendance_rows.append(
        {
            "batch_id": batch_id,
            "employee_code": attendance.employee_code,
            "employee_name": attendance.employee_name,
            "location": attendance.location,
            "date": attendance.date,
            "status": attendance.status,
            "raw_status": attendance.raw_status,
            "in_time": attendance.in_time,
            "out_time": attendance.out_time,
            "total_minutes": attendance.total_minutes,
        }
    )


def finalize_batch(batch, status, error=None):
    batch.status = status
    batch.error = error
    _write_batch_outputs(batch)
    LOGGER.info(
        "Batch %s %s: %s attendance rows, %s failures, %s data-quality issues",
        batch.id,
        status,
        len(batch.attendance_rows),
        len(batch.failures),
        len(batch.data_quality_issues),
    )
    if error:
        LOGGER.error("Batch %s error: %s", batch.id, error)


def _batch_summary(batch):
    return {
        "batch_id": batch.id,
        "file_name": batch.file_name,
        "file_hash": batch.file_hash,
        "location": batch.location,
        "status": batch.status,
        "error": batch.error,
        "period": batch.period,
        "employees_failed": len(batch.failures),
        "data_quality_issues": len(batch.data_quality_issues),
        "attendance_upserted": len(batch.attendance_rows),
        "failures": batch.failures,
        "issues": batch.data_quality_issues,
        "api_result": batch.api_result,
    }


def _write_batch_outputs(batch):
    output_dir = Path(CONFIG["output_dir"])
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / f"batch_{batch.id}_summary.json").write_text(
        json.dumps(_batch_summary(batch), indent=2), encoding="utf-8"
    )
    (output_dir / f"batch_{batch.id}_attendance.json").write_text(
        json.dumps(batch.attendance_rows, indent=2), encoding="utf-8"
    )
    LOGGER.info(
        "Wrote attendance JSON to %s",
        output_dir / f"batch_{batch.id}_attendance.json",
    )
    _prune_output_dir(output_dir)


# Runtime files that must stay in output/
_OUTPUT_KEEP_NAMES = {
    ".gitkeep",
    "processed_hashes.json",
    "importer.lock",
    "importer.log",
    "last_run.json",
    "last_run.txt",
    "scheduler_runs.log",
}

# One-off handoff / sample junk left from earlier debugging (safe to delete)
_OUTPUT_JUNK_PREFIXES = (
    "API_HANDOFF_",
    "WTT_ATTENDANCE_",
    "wtt_attendance_",
    "sample_",
)
_BATCH_FILE_RE = re.compile(r"^batch_([0-9a-fA-F]+)_((?:attendance)|(?:summary))\.json$")


def _prune_output_dir(output_dir: Path | None = None, keep: int | None = None) -> None:
    """
    Keep only the newest `keep` import batches in output/, and delete leftover
    handoff/sample files that are not needed at runtime.
    """
    output_dir = Path(output_dir or CONFIG["output_dir"])
    keep = int(CONFIG.get("output_batch_keep", 10) if keep is None else keep)
    if keep < 1 or not output_dir.is_dir():
        return

    # 1) Drop unused handoff / sample artifacts
    junk_names = {
        "API_HANDOFF_IMPORTER_TO_RECEIVER.txt",
        "WTT_ATTENDANCE_INSERT_PROMPT.txt",
        "wtt_attendance_handoff_sample.json",
    }
    for path in list(output_dir.iterdir()):
        if not path.is_file():
            continue
        name = path.name
        if name in _OUTPUT_KEEP_NAMES or _BATCH_FILE_RE.match(name):
            continue
        is_junk = (
            name in junk_names
            or name.startswith(_OUTPUT_JUNK_PREFIXES)
            or (name.lower().endswith(".sql") and name.lower().startswith("sample_"))
        )
        if not is_junk:
            continue
        try:
            path.unlink()
            LOGGER.info("Removed unused output file: %s", name)
        except Exception as exc:
            LOGGER.warning("Could not remove %s: %s", path, exc)

    # 2) Keep newest N batch ids (by attendance.json mtime, else summary)
    batches: dict[str, dict] = {}
    for path in output_dir.iterdir():
        if not path.is_file():
            continue
        match = _BATCH_FILE_RE.match(path.name)
        if not match:
            continue
        batch_id, kind = match.group(1), match.group(2)
        info = batches.setdefault(batch_id, {"files": [], "mtime": 0.0})
        info["files"].append(path)
        try:
            mtime = path.stat().st_mtime
        except OSError:
            continue
        # Prefer attendance.json mtime as the batch timestamp when present.
        if kind == "attendance":
            info["mtime"] = mtime
            info["has_attendance"] = True
        elif not info.get("has_attendance"):
            info["mtime"] = max(info["mtime"], mtime)

    if len(batches) <= keep:
        return

    ordered = sorted(batches.items(), key=lambda item: item[1]["mtime"], reverse=True)
    for batch_id, info in ordered[keep:]:
        for path in info["files"]:
            try:
                path.unlink()
                LOGGER.info("Pruned old batch output (keep %s): %s", keep, path.name)
            except Exception as exc:
                LOGGER.warning("Could not prune %s: %s", path, exc)


def _row_to_api_dict(row) -> dict:
    if isinstance(row, AttendanceRecord):
        return row.to_api_dict()
    return {
        "employee_code": row["employee_code"],
        "employee_name": row["employee_name"],
        "attendance_date": row["date"],
        "status": row.get("raw_status") or row.get("status"),
        "in_time": row.get("in_time"),
        "out_time": row.get("out_time"),
        "total_minutes": row.get("total_minutes"),
    }


def send_attendance_to_ddo_api(batch):
    endpoint = os.getenv("DDO_API_ENDPOINT", CONFIG.get("api_endpoint", "")).strip()
    token = (
        os.getenv("ATTENDANCE_INTEGRATION_TOKEN")
        or os.getenv("DDO_API_TOKEN", CONFIG.get("api_token", ""))
        or ""
    ).strip()
    location_code = validate_location(batch.location or get_location_from_config())
    timeout = CONFIG["api_timeout_seconds"]
    max_retries = max(1, CONFIG["api_max_retries"])
    chunk_size = max(1, CONFIG.get("api_chunk_size", 50))
    chunk_delay = max(0.0, float(CONFIG.get("api_chunk_delay_seconds", 2)))
    max_body_bytes = max(20_000, CONFIG.get("api_max_body_bytes", 90_000))

    if not endpoint:
        raise RuntimeError("DDO_API_ENDPOINT is not configured")
    if "/admin/" in endpoint.lower():
        raise RuntimeError(
            "DDO_API_ENDPOINT is the NocoBase admin page, not the import API. "
            "Use https://<host>/api/attendance/import"
        )
    if not _is_allowed_endpoint(endpoint):
        raise RuntimeError(
            f"DDO++ API endpoint must use HTTPS (http is allowed only for localhost): {endpoint}"
        )
    if not token:
        raise RuntimeError("ATTENDANCE_INTEGRATION_TOKEN or DDO_API_TOKEN is not configured")
    if not location_code:
        raise RuntimeError("LOCATION_CODE is not configured")

    all_records = [_row_to_api_dict(row) for row in batch.attendance_rows]
    if not all_records:
        # Still notify API with empty records (server may reject; keep previous behavior).
        chunks = [[]]
    else:
        chunks = _split_records_for_api(all_records, chunk_size, max_body_bytes, location_code, batch)

    LOGGER.info(
        "POST %s location_code=%s batch_id=%s records=%s chunks=%s chunk_size=%s delay=%ss",
        endpoint,
        location_code,
        batch.id,
        len(all_records),
        len(chunks),
        chunk_size,
        chunk_delay,
    )

    chunk_results = []
    hard_error = None
    for chunk_index, records in enumerate(chunks, start=1):
        if chunk_index > 1 and chunk_delay > 0:
            LOGGER.info(
                "Waiting %.1fs before chunk %s/%s so the database can settle",
                chunk_delay,
                chunk_index,
                len(chunks),
            )
            time.sleep(chunk_delay)
        chunk_id = batch.id if len(chunks) == 1 else f"{batch.id}-p{chunk_index}"
        try:
            result = _post_attendance_chunk(
                endpoint=endpoint,
                token=token,
                location_code=location_code,
                batch=batch,
                records=records,
                chunk_id=chunk_id,
                chunk_index=chunk_index,
                chunk_total=len(chunks),
                timeout=timeout,
                max_retries=max_retries,
                max_body_bytes=max_body_bytes,
            )
            chunk_results.append(result)
        except RuntimeError as exc:
            # Catastrophic (5xx / network / timeout after retries): stop remaining chunks.
            hard_error = str(exc)
            chunk_results.append(
                {
                    "ok": False,
                    "outcome": "ERROR",
                    "batch_id": chunk_id,
                    "records": len(records),
                    "error": hard_error,
                }
            )
            LOGGER.error(
                "Stopping chunk upload after catastrophic failure on chunk %s/%s: %s",
                chunk_index,
                len(chunks),
                hard_error,
            )
            break

    flat = _flatten_chunk_results(chunk_results)
    outcomes = [str(item.get("outcome") or "").upper() for item in flat]
    partial_count = sum(1 for o in outcomes if o == "PARTIAL")
    failed_count = sum(1 for o in outcomes if o == "FAILED")
    success_count = sum(1 for o in outcomes if o == "SUCCESS")
    error_count = sum(1 for o in outcomes if o == "ERROR")

    if hard_error or error_count:
        batch.api_result = {
            "ok": False,
            "chunks": len(chunks),
            "records": len(all_records),
            "responses": chunk_results,
            "summary": {
                "success": success_count,
                "partial": partial_count,
                "failed": failed_count,
                "error": error_count,
            },
            "error": hard_error or "One or more chunks failed catastrophically",
        }
        _write_batch_outputs(batch)
        raise RuntimeError(batch.api_result["error"])

    # All chunks attempted: fail the file only if every chunk was a hard business FAILED.
    if flat and failed_count == len(flat):
        msg = (
            f"All {failed_count} API chunk(s) returned FAILED; "
            "no attendance rows were accepted"
        )
        batch.api_result = {
            "ok": False,
            "chunks": len(chunks),
            "records": len(all_records),
            "responses": chunk_results,
            "summary": {
                "success": success_count,
                "partial": partial_count,
                "failed": failed_count,
                "error": error_count,
            },
            "error": msg,
        }
        _write_batch_outputs(batch)
        raise RuntimeError(msg)

    overall_ok = True
    if partial_count or failed_count:
        LOGGER.warning(
            "Chunk upload finished with partial/row failures: "
            "success=%s partial=%s failed=%s (remaining chunks were still sent)",
            success_count,
            partial_count,
            failed_count,
        )

    batch.api_result = {
        "ok": overall_ok,
        "partial": partial_count > 0 or failed_count > 0,
        "chunks": len(chunks),
        "records": len(all_records),
        "responses": chunk_results,
        "summary": {
            "success": success_count,
            "partial": partial_count,
            "failed": failed_count,
            "error": error_count,
        },
    }
    _write_batch_outputs(batch)
    return batch.api_result


def _flatten_chunk_results(results):
    """Flatten nested split-chunk responses into leaf chunk results."""
    flat = []
    for item in results or []:
        if not isinstance(item, dict):
            continue
        parts = item.get("parts")
        if parts:
            flat.extend(_flatten_chunk_results(parts))
        else:
            flat.append(item)
    return flat


def _split_records_for_api(records, chunk_size, max_body_bytes, location_code, batch):
    """Split records so each JSON body stays under the proxy body limit."""
    chunks = []
    current = []
    for record in records:
        candidate = current + [record]
        if len(candidate) > chunk_size or (
            current and _payload_size(location_code, batch, candidate) > max_body_bytes
        ):
            chunks.append(current)
            current = [record]
        else:
            current = candidate
    if current or not chunks:
        chunks.append(current)
    return chunks


def _payload_size(location_code, batch, records):
    payload = {
        "location_code": location_code,
        "report_from": batch.period.get("from") or batch.period.get("start"),
        "report_to": batch.period.get("to") or batch.period.get("end"),
        "source_file": batch.file_name,
        "file_hash": batch.file_hash,
        "batch_id": batch.id,
        "records": records,
    }
    return len(json.dumps(payload).encode("utf-8"))


def _post_attendance_chunk(
    endpoint,
    token,
    location_code,
    batch,
    records,
    chunk_id,
    chunk_index,
    chunk_total,
    timeout,
    max_retries,
    max_body_bytes,
):
    # If a single chunk is still too large, keep splitting until it fits or one record left.
    working = list(records)
    while len(working) > 1 and _payload_size(location_code, batch, working) > max_body_bytes:
        mid = max(1, len(working) // 2)
        LOGGER.warning(
            "Chunk %s/%s still too large (%s bytes); splitting %s -> %s + %s",
            chunk_index,
            chunk_total,
            _payload_size(location_code, batch, working),
            len(working),
            mid,
            len(working) - mid,
        )
        first = _post_attendance_chunk(
            endpoint,
            token,
            location_code,
            batch,
            working[:mid],
            f"{chunk_id}a",
            chunk_index,
            chunk_total,
            timeout,
            max_retries,
            max_body_bytes,
        )
        second = _post_attendance_chunk(
            endpoint,
            token,
            location_code,
            batch,
            working[mid:],
            f"{chunk_id}b",
            chunk_index,
            chunk_total,
            timeout,
            max_retries,
            max_body_bytes,
        )
        return {"ok": True, "split": True, "parts": [first, second]}

    payload = {
        "location_code": location_code,
        "report_from": batch.period.get("from") or batch.period.get("start"),
        "report_to": batch.period.get("to") or batch.period.get("end"),
        "source_file": batch.file_name,
        "file_hash": batch.file_hash,
        "batch_id": chunk_id,
        "records": working,
    }
    body = json.dumps(payload).encode("utf-8")
    LOGGER.info(
        "POST chunk %s/%s batch_id=%s records=%s bytes=%s",
        chunk_index,
        chunk_total,
        chunk_id,
        len(working),
        len(body),
    )

    last_error = None
    last_status_code = None
    last_parsed = {}
    for attempt in range(1, max_retries + 1):
        try:
            status_code, response_text = _post_json(endpoint, token, location_code, body, timeout)
            parsed = _parse_api_response(response_text)
            last_status_code = status_code
            last_parsed = parsed

            if 200 <= status_code < 300 and _api_accepted_business(parsed):
                LOGGER.info(
                    "DDO++ API accepted chunk %s/%s (HTTP %s, attempt %s/%s)",
                    chunk_index,
                    chunk_total,
                    status_code,
                    attempt,
                    max_retries,
                )
                return {
                    "ok": True,
                    "outcome": "SUCCESS",
                    "status_code": status_code,
                    "attempt": attempt,
                    "batch_id": chunk_id,
                    "records": len(working),
                    "response": parsed,
                }

            # HTTP 207 / body status PARTIAL: valid rows already saved — continue upload loop.
            if _api_is_partial(parsed, status_code):
                detail = _format_api_failure(parsed, status_code, response_text)
                LOGGER.warning(
                    "DDO++ API partial success for chunk %s/%s (HTTP %s): %s "
                    "(continuing remaining chunks)",
                    chunk_index,
                    chunk_total,
                    status_code,
                    detail,
                )
                return {
                    "ok": True,
                    "outcome": "PARTIAL",
                    "status_code": status_code,
                    "attempt": attempt,
                    "batch_id": chunk_id,
                    "records": len(working),
                    "response": parsed,
                    "warning": detail,
                }

            # Business FAILED (typically HTTP 422): do not retry; return so caller can continue.
            if _api_is_business_failed(parsed, status_code):
                last_error = _format_api_failure(parsed, status_code, response_text)
                LOGGER.error(
                    "DDO++ API business FAILED for chunk %s/%s (HTTP %s): %s "
                    "(continuing remaining chunks)",
                    chunk_index,
                    chunk_total,
                    status_code,
                    last_error,
                )
                return {
                    "ok": False,
                    "outcome": "FAILED",
                    "status_code": status_code,
                    "attempt": attempt,
                    "batch_id": chunk_id,
                    "records": len(working),
                    "response": parsed,
                    "error": last_error,
                }

            last_error = _format_api_failure(parsed, status_code, response_text)
            LOGGER.error(
                "DDO++ API rejected chunk %s/%s (HTTP %s, attempt %s/%s): %s",
                chunk_index,
                chunk_total,
                status_code,
                attempt,
                max_retries,
                last_error,
            )
            # 413: shrink chunk and retry immediately instead of giving up.
            if status_code == 413 and len(working) > 1:
                mid = max(1, len(working) // 2)
                LOGGER.warning("HTTP 413 — retrying as two smaller chunks (%s + %s)", mid, len(working) - mid)
                first = _post_attendance_chunk(
                    endpoint, token, location_code, batch, working[:mid], f"{chunk_id}a",
                    chunk_index, chunk_total, timeout, max_retries, max_body_bytes,
                )
                second = _post_attendance_chunk(
                    endpoint, token, location_code, batch, working[mid:], f"{chunk_id}b",
                    chunk_index, chunk_total, timeout, max_retries, max_body_bytes,
                )
                return {"ok": True, "split_after_413": True, "parts": [first, second]}
            # Other 4xx (auth, validation): do not retry; treat as catastrophic below.
            if status_code < 500 and status_code != 429:
                break
        except urllib.error.HTTPError as exc:
            response_text = _read_http_error_body(exc)
            parsed = _parse_api_response(response_text)
            last_status_code = exc.code
            last_parsed = parsed
            last_error = _format_api_failure(parsed, exc.code, response_text)

            if _api_is_partial(parsed, exc.code):
                LOGGER.warning(
                    "DDO++ API partial success for chunk %s/%s (HTTP %s): %s "
                    "(continuing remaining chunks)",
                    chunk_index,
                    chunk_total,
                    exc.code,
                    last_error,
                )
                return {
                    "ok": True,
                    "outcome": "PARTIAL",
                    "status_code": exc.code,
                    "attempt": attempt,
                    "batch_id": chunk_id,
                    "records": len(working),
                    "response": parsed,
                    "warning": last_error,
                }

            if _api_is_business_failed(parsed, exc.code):
                LOGGER.error(
                    "DDO++ API business FAILED for chunk %s/%s (HTTP %s): %s "
                    "(continuing remaining chunks)",
                    chunk_index,
                    chunk_total,
                    exc.code,
                    last_error,
                )
                return {
                    "ok": False,
                    "outcome": "FAILED",
                    "status_code": exc.code,
                    "attempt": attempt,
                    "batch_id": chunk_id,
                    "records": len(working),
                    "response": parsed,
                    "error": last_error,
                }

            LOGGER.error(
                "DDO++ API HTTP error for chunk %s/%s (HTTP %s, attempt %s/%s): %s",
                chunk_index,
                chunk_total,
                exc.code,
                attempt,
                max_retries,
                last_error,
            )
            if exc.code == 413 and len(working) > 1:
                mid = max(1, len(working) // 2)
                LOGGER.warning("HTTP 413 — retrying as two smaller chunks (%s + %s)", mid, len(working) - mid)
                first = _post_attendance_chunk(
                    endpoint, token, location_code, batch, working[:mid], f"{chunk_id}a",
                    chunk_index, chunk_total, timeout, max_retries, max_body_bytes,
                )
                second = _post_attendance_chunk(
                    endpoint, token, location_code, batch, working[mid:], f"{chunk_id}b",
                    chunk_index, chunk_total, timeout, max_retries, max_body_bytes,
                )
                return {"ok": True, "split_after_413": True, "parts": [first, second]}
            if exc.code < 500 and exc.code != 429:
                break
        except urllib.error.URLError as exc:
            last_error = f"Network error: {exc.reason}"
            LOGGER.error(
                "DDO++ API network error for chunk %s/%s (attempt %s/%s): %s",
                chunk_index,
                chunk_total,
                attempt,
                max_retries,
                exc.reason,
            )
        except ssl.SSLError as exc:
            last_error = f"TLS error: {exc}"
            LOGGER.error("DDO++ API TLS error for chunk %s/%s: %s", chunk_index, chunk_total, exc)
            break
        except TimeoutError:
            last_error = "Timeout"
            LOGGER.error(
                "DDO++ API timeout for chunk %s/%s (attempt %s/%s)",
                chunk_index,
                chunk_total,
                attempt,
                max_retries,
            )

        if attempt < max_retries:
            delay = min(2 ** (attempt - 1), 8)
            LOGGER.info("Retrying DDO++ API POST in %s seconds", delay)
            time.sleep(delay)

    batch.api_result = {
        "ok": False,
        "error": last_error,
        "failed_chunk": chunk_id,
        "status_code": last_status_code,
        "response": last_parsed,
    }
    _write_batch_outputs(batch)
    raise RuntimeError(f"Failed to send attendance JSON to DDO++ API: {last_error}")


def _is_allowed_endpoint(endpoint):
    parsed = urllib.parse.urlparse(endpoint)
    host = (parsed.hostname or "").lower()
    path = (parsed.path or "").lower()
    if "/admin/" in path:
        return False
    if parsed.scheme == "https" and parsed.netloc:
        return True
    return parsed.scheme == "http" and host in {"127.0.0.1", "localhost", "::1"}


def _ssl_context():
    context = ssl.create_default_context()
    ca_file = os.getenv("DDO_API_CA_FILE", "").strip()
    if ca_file:
        context.load_verify_locations(cafile=ca_file)
    return context


def _post_json(endpoint, token, location_code, body, timeout):
    request = urllib.request.Request(
        endpoint,
        data=body,
        method="POST",
        headers={
            "Content-Type": "application/json",
            "Accept": "application/json",
            "Authorization": f"Bearer {token}",
            "X-Location-Code": location_code,
        },
    )
    kwargs = {"timeout": timeout}
    if urllib.parse.urlparse(endpoint).scheme == "https":
        kwargs["context"] = _ssl_context()
    with urllib.request.urlopen(request, **kwargs) as response:
        response_text = response.read().decode("utf-8", errors="replace")
        return response.status, response_text


def _read_http_error_body(exc):
    try:
        return exc.read().decode("utf-8", errors="replace")
    except Exception:
        return str(exc)


def _parse_api_response(response_text):
    if not response_text:
        return {}
    try:
        return json.loads(response_text)
    except json.JSONDecodeError:
        return {"raw": _clip(response_text)}


def _api_business_status(parsed) -> str:
    """Normalize API body status (SUCCESS / PARTIAL / FAILED / ...)."""
    if not isinstance(parsed, dict):
        return ""
    return str(parsed.get("status") or "").strip().upper()


def _api_accepted_business(parsed) -> bool:
    """
    True only when the API fully accepted the chunk (SUCCESS).

    PARTIAL (HTTP 207) and FAILED (HTTP 422) are handled separately so the
    upload loop can continue sending remaining chunks.
    """
    if isinstance(parsed, dict) and parsed.get("ok") is False:
        return False
    status = _api_business_status(parsed)
    if status in {"FAILED", "PARTIAL"}:
        return False
    return True


def _api_is_partial(parsed, status_code=None) -> bool:
    """True for HTTP 207 Multi-Status or body status PARTIAL (valid rows saved)."""
    if status_code == 207:
        return True
    return _api_business_status(parsed) == "PARTIAL"


def _api_is_business_failed(parsed, status_code=None) -> bool:
    """True for body status FAILED or HTTP 422 Unprocessable Entity."""
    if status_code == 422:
        return True
    return _api_business_status(parsed) == "FAILED"


def _format_api_failure(parsed, status_code, response_text) -> str:
    status = _api_business_status(parsed) or f"HTTP_{status_code}"
    parts = [f"status={status}", f"http={status_code}"]
    if isinstance(parsed, dict):
        for key in ("failed", "inserted", "updated", "processed", "batch_id"):
            if key in parsed and parsed[key] is not None:
                parts.append(f"{key}={parsed[key]}")
        details = parsed.get("failed_details") or parsed.get("failures") or []
        if isinstance(details, list) and details:
            samples = []
            for row in details[:5]:
                if isinstance(row, dict):
                    code = row.get("employee_code") or "?"
                    err = row.get("error") or row.get("reason") or row
                    samples.append(f"{code}: {err}")
                else:
                    samples.append(str(row))
            parts.append("examples=[" + "; ".join(samples) + "]")
            if len(details) > 5:
                parts.append(f"(+{len(details) - 5} more)")
    else:
        parts.append(_clip(response_text))
    return " | ".join(parts)


def _clip(text, limit=500):
    value = str(text or "").replace("\n", " ").strip()
    if len(value) <= limit:
        return value
    return value[:limit] + "..."


def move_to_processed(file):
    _move_file(file, CONFIG["processed_dir"])
    _prune_folder_keep_newest(CONFIG["processed_dir"], CONFIG.get("archive_keep", 15))


def move_to_failed(file):
    _move_file(file, CONFIG["failed_dir"])
    _prune_folder_keep_newest(CONFIG["failed_dir"], CONFIG.get("archive_keep", 15))


def _prune_folder_keep_newest(folder, keep: int = 15) -> None:
    """Keep only the newest `keep` files in processed/ or failed/."""
    path = Path(folder)
    keep = int(keep)
    if keep < 1 or not path.is_dir():
        return
    files = [p for p in path.iterdir() if p.is_file() and p.name != ".gitkeep"]
    if len(files) <= keep:
        return
    files.sort(key=lambda p: (p.stat().st_mtime, p.name), reverse=True)
    for old in files[keep:]:
        try:
            old.unlink()
            LOGGER.info("Pruned old archive (keep %s) in %s: %s", keep, path.name, old.name)
        except Exception as exc:
            LOGGER.warning("Could not prune %s: %s", old, exc)


def _row_label(row):
    for cell in row[:6]:
        value = re.sub(r"[\s._-]", "", _cell_text(cell).lower())
        if value in {"status", "intime", "outtime", "total"}:
            return value
    return None


def _is_department_row(row, header_offset):
    name = _cell_text(row[header_offset] if header_offset < len(row) else "")
    if not name:
        name = next((cell for cell in row if _cell_text(cell)), "")
    lowered = name.lower()
    if lowered in DEPARTMENT_NAMES:
        return True
    return bool(re.match(r"^department\b", lowered))


def _unmapped_code(name):
    slug = re.sub(r"[^a-z0-9]+", "_", (name or "").lower()).strip("_")
    return f"UNMAPPED:{slug}" if slug else ""


def _day_values(row, day_columns):
    return [_cell_text(row[idx]) if idx < len(row) else "" for idx in day_columns]


def _parse_time(value):
    text = _cell_text(value)
    match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", text)
    if not match:
        return None
    return f"{int(match.group(1)):02d}:{match.group(2)}" + (f":{match.group(3)}" if match.group(3) else "")


def _duration_to_minutes(value):
    text = _cell_text(value)
    match = re.fullmatch(r"(\d{1,3}):(\d{2})(?::(\d{2}))?", text)
    if not match:
        return None
    return int(match.group(1)) * 60 + int(match.group(2))


def _identity_from_labeled_row(row):
    cells = [_cell_text(cell) for cell in row]
    if not any(re.match(r"emp\.?\s*code", cell, re.I) for cell in cells):
        return None
    code = ""
    name = ""
    for idx, cell in enumerate(cells):
        if re.match(r"emp\.?\s*code", cell, re.I):
            for nxt in cells[idx + 1 :]:
                if nxt and not re.match(r"emp\.?\s*name", nxt, re.I):
                    code = nxt
                    break
        if re.match(r"emp\.?\s*name", cell, re.I):
            for nxt in cells[idx + 1 :]:
                if nxt:
                    name = nxt
                    break
    return code, name


def _days_from_grouped(period, grouped):
    statuses = grouped.get("status") or [""] * len(period["dates"])
    in_times = grouped.get("intime") or [""] * len(period["dates"])
    out_times = grouped.get("outtime") or [""] * len(period["dates"])
    totals = grouped.get("total") or [""] * len(period["dates"])
    days = []
    for idx, day_date in enumerate(period["dates"]):
        raw = statuses[idx] if idx < len(statuses) else ""
        days.append(
            DayRecord(
                date=day_date,
                status=_normalize_status(raw),
                raw_status=raw,
                header=period["day_headers"][idx],
                in_time=_parse_time(in_times[idx] if idx < len(in_times) else ""),
                out_time=_parse_time(out_times[idx] if idx < len(out_times) else ""),
                total_minutes=_duration_to_minutes(totals[idx] if idx < len(totals) else ""),
            )
        )
    return days


def _extract_punch_blocks(rows, period):
    blocks = []
    idx = period["header_row_index"] + 1
    while idx < len(rows):
        row = rows[idx]
        if not any(row) or _is_department_row(row, period["header_offset"]) or _row_label(row):
            idx += 1
            continue
        identity = _identity_from_labeled_row(row)
        grouped = {}
        if identity:
            code, name = identity
            next_idx = idx + 1
            while next_idx < len(rows) and _row_label(rows[next_idx]):
                grouped[_row_label(rows[next_idx])] = _day_values(rows[next_idx], period["day_columns"])
                next_idx += 1
            if not grouped.get("status"):
                grouped["status"] = _day_values(row, period["day_columns"])
            idx = next_idx
        else:
            first = _cell_text(row[0] if row else "")
            name = _cell_text(row[period["header_offset"]] if period["header_offset"] < len(row) else "") or _cell_text(row[1] if len(row) > 1 else "")
            code = first if re.fullmatch(r"\d{3,}", first) else ""
            if not code and not name:
                idx += 1
                continue
            grouped["status"] = _day_values(row, period["day_columns"])
            next_idx = idx + 1
            while next_idx < len(rows) and _row_label(rows[next_idx]):
                grouped[_row_label(rows[next_idx])] = _day_values(rows[next_idx], period["day_columns"])
                next_idx += 1
            idx = next_idx
        if not code and not name:
            continue
        blocks.append(EmployeeBlock(code=code, name=name, days=_days_from_grouped(period, grouped)))
    return blocks


def _extract_grid_blocks(rows, period):
    blocks = []
    has_code_header = period.get("has_code_header")
    for row in rows[period["header_row_index"] + 1 :]:
        if not any(row) or _is_department_row(row, period["header_offset"]) or _row_label(row):
            continue
        first = _cell_text(row[0] if row else "")
        second = _cell_text(row[1] if len(row) > 1 else "")
        if has_code_header:
            code = first
            name = second or _cell_text(row[period["header_offset"]] if period["header_offset"] < len(row) else "")
        else:
            name = first or _cell_text(row[period["header_offset"]] if period["header_offset"] < len(row) else "")
            code = first if re.fullmatch(r"\d{3,}", first) else ""
        if re.match(r"emp\.?\s*code", first, re.I) or first.lower() in {"name", "days"}:
            continue
        if not code and not name:
            continue
        grouped = {"status": _day_values(row, period["day_columns"])}
        blocks.append(EmployeeBlock(code=code, name=name, days=_days_from_grouped(period, grouped)))
    return blocks


def process_file(file, location=None, raise_on_error=True):
    _configure_logging()
    resolved_location = validate_location(location or get_location_from_config())
    os.environ["LOCATION_CODE"] = resolved_location

    batch = create_import_batch(file, resolved_location)

    try:
        report = read_excel(file)
        extract_report_period(report)

        for employee_block in extract_employee_blocks(report):
            employee = resolve_employee(employee_block.code, resolved_location)

            if not employee:
                record_failure(batch, employee_block.code, "Unknown employee")
                continue

            for day in employee_block.days:
                if not day.status:
                    record_data_quality_issue(batch, employee.code, day.date)
                    continue

                attendance = transform(day, employee, resolved_location)
                upsert_attendance(attendance, batch.id)

        finalize_batch(batch, "SUCCESS")
        send_attendance_to_ddo_api(batch)
        move_to_processed(file)
        return batch

    except Exception as exc:
        LOGGER.error("Import failed for %s: %s", Path(file).name, exc)
        finalize_batch(batch, "FAILED", str(exc))
        move_to_failed(file)
        if raise_on_error:
            raise
        return batch


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.date() in (date(1899, 12, 30), date(1899, 12, 31), date(1900, 1, 1)):
            return value.strftime("%H:%M:%S")
        if value.time() == datetime.min.time():
            return value.date().isoformat()
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, dt_time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, timedelta):
        total_seconds = int(value.total_seconds())
        hours, remainder = divmod(abs(total_seconds), 3600)
        minutes, seconds = divmod(remainder, 60)
        sign = "-" if total_seconds < 0 else ""
        if seconds:
            return f"{sign}{hours}:{minutes:02d}:{seconds:02d}"
        return f"{sign}{hours}:{minutes:02d}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _select_sheet(sheets):
    if not sheets:
        raise ValueError("Workbook has no sheets")
    best_sheet = sheets[0]
    best_score = -1
    for sheet in sheets:
        preview = []
        for row in sheet["rows"][:12]:
            preview.extend(row[:40])
        joined = " ".join(preview).lower()
        score = 0
        if "basicworkduration" in sheet["name"].lower():
            score += 50
        if "emp. code" in joined or "emp code" in joined:
            score += 20
        if any(_parse_day_header(cell) for cell in preview):
            score += 15
        if "status" in joined and "intime" in joined.replace(" ", ""):
            score += 15
        score += min(len(sheet["rows"]), 200) / 10
        if score > best_score:
            best_score = score
            best_sheet = sheet
    return best_sheet


def _day_headers(headers):
    return [header for header in headers if _parse_day_header(header)]


def _header_offset(headers):
    offset = 0
    while offset < len(headers) and not _parse_day_header(headers[offset]):
        offset += 1
    return max(offset - 1, 0)


def _parse_day_header(header):
    match = re.fullmatch(r"(\d{1,2})\s*([A-Za-z]{1,2})", str(header).strip())
    if not match:
        return None
    weekday = match.group(2).upper()
    if weekday not in WEEKDAY_CODES:
        return None
    return int(match.group(1)), weekday


def _infer_year_month(file_name, report_text=""):
    months = {
        "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
        "july": 7, "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
        "jan": 1, "feb": 2, "mar": 3, "apr": 4, "jun": 6, "jul": 7, "aug": 8,
        "sep": 9, "oct": 10, "nov": 11, "dec": 12,
    }
    month = date.today().month
    haystack = f"{file_name} {report_text}"
    for name, number in months.items():
        if re.search(rf"\b{name}\b", haystack, flags=re.IGNORECASE):
            month = number
            break

    year = date.today().year
    file_year = re.search(r"(20\d{2})", file_name)
    period_year = re.search(
        r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|"
        r"aug(?:ust)?|sep(?:tember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
        r"[a-z]*\s+\d{1,2}\s+(20\d{2})"
        r"|(20\d{2})\s+To\s+(20\d{2})"
        r"|To\s+[A-Za-z]{3,9}\s+\d{1,2}\s+(20\d{2})",
        report_text,
        flags=re.IGNORECASE,
    )
    if file_year:
        year = int(file_year.group(1))
    elif period_year:
        year = int(next(g for g in period_year.groups() if g))
    return year, month


def _map_headers_to_dates(day_headers, year, month):
    """Map eSSL day headers like '23 T' to calendar dates in (year, month).

    Older logic did ``month - 1`` when the first day number was > 20 (meant for
    rare cross-month payroll grids). That broke daily single-column exports:
    September 23 became August 23. We now keep the inferred report month and
    only roll forward when day numbers decrease (e.g. 30 then 1).
    """
    parsed = [_parse_day_header(header) for header in day_headers]
    parsed = [item for item in parsed if item]
    if not parsed:
        return []

    dates = []
    current_year, current_month = year, month
    last_day = None
    for day_number, _weekday in parsed:
        if last_day is not None and day_number < last_day:
            current_month += 1
            if current_month > 12:
                current_month = 1
                current_year += 1
        dates.append(date(current_year, current_month, day_number))
        last_day = day_number
    return dates


def _normalize_status(raw):
    cleaned = re.sub(r"[^A-Za-z]", "", raw or "").upper()
    return cleaned


def _move_file(file, target_dir):
    source = Path(file)
    destination_dir = Path(target_dir)
    destination_dir.mkdir(parents=True, exist_ok=True)
    destination = destination_dir / source.name
    if destination.exists():
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        destination = destination_dir / f"{source.stem}_{stamp}{source.suffix}"
    shutil.move(str(source), str(destination))
    LOGGER.info("Moved %s -> %s", source.name, destination)


def _default_input_file():
    candidates = [
        REPO_ROOT / "NOIDA attendance report august.xls",
        Path(r"C:\Users\WTT\Downloads\NOIDA attendance report august.xls"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def infer_location(file_path, explicit=None):
    if explicit:
        return normalize_location(explicit)
    path = Path(file_path)
    parent = path.parent.name.upper()
    if parent in ALLOWED_LOCATIONS:
        return ALLOWED_LOCATIONS[parent]
    haystack = path.name.upper()
    for alias, code in ALLOWED_LOCATIONS.items():
        if re.search(rf"\b{re.escape(alias)}\b", haystack):
            return code
    return None


def _office_drop_dir() -> Path | None:
    """Single per-PC folder where HR drops Excel. Location comes from LOCATION_CODE."""
    raw = (CONFIG.get("attendance_dir") or "").strip()
    if not raw:
        return None
    path = Path(raw)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _is_excel_report(path: Path) -> bool:
    return path.is_file() and path.suffix.lower() in EXCEL_SUFFIXES and not path.name.startswith("~$")

def _is_file_ready(path: Path) -> bool:
    try:
        age = time.time() - path.stat().st_mtime
    except OSError:
        return False
    if age < CONFIG["stable_seconds"]:
        LOGGER.info("Skipping %s; still being written (%ss old)", path.name, int(age))
        return False
    try:
        with open(path, "rb") as handle:
            if sys.platform == "win32":
                import msvcrt

                msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
        return True
    except OSError:
        LOGGER.info("Skipping %s; file is locked by another process", path.name)
        return False


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _hash_store_path() -> Path:
    return Path(CONFIG["output_dir"]) / "processed_hashes.json"


def _load_processed_hashes() -> dict:
    store = _hash_store_path()
    if not store.exists():
        return {}
    try:
        return json.loads(store.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}


def _remember_processed_hash(file_hash: str, file_name: str, location: str) -> None:
    store = _hash_store_path()
    store.parent.mkdir(parents=True, exist_ok=True)
    payload = _load_processed_hashes()
    payload[file_hash] = {
        "file_name": file_name,
        "location": location,
        "processed_at": datetime.now().isoformat(timespec="seconds"),
    }
    store.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def _lock_path() -> Path:
    return Path(CONFIG["output_dir"]) / "importer.lock"


class _SingleInstance:
    def __init__(self, path: Path):
        self.path = path
        self.handle = None

    def acquire(self) -> bool:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.handle = open(self.path, "a+b")
        try:
            if self.handle.tell() == 0:
                self.handle.write(b"0")
                self.handle.flush()
            self.handle.seek(0)
            if sys.platform == "win32":
                import msvcrt

                msvcrt.locking(self.handle.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl

                fcntl.flock(self.handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError:
            self.handle.close()
            self.handle = None
            return False
        self.handle.seek(0)
        self.handle.truncate()
        stamp = f"{os.getpid()} {datetime.now().isoformat(timespec='seconds')}\n"
        self.handle.write(stamp.encode("utf-8"))
        self.handle.flush()
        return True

    def release(self) -> None:
        if self.handle is None:
            return
        try:
            if sys.platform == "win32":
                import msvcrt

                self.handle.seek(0)
                msvcrt.locking(self.handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl

                fcntl.flock(self.handle.fileno(), fcntl.LOCK_UN)
        except OSError:
            pass
        self.handle.close()
        self.handle = None


def _scan_flat_drop(root: Path, office_location: str, seen: set) -> list:
    """Pick Excel files directly in root; tag them with this PC's LOCATION_CODE."""
    found = []
    try:
        children = sorted(root.iterdir())
    except OSError as exc:
        LOGGER.warning("Could not read drop folder %s: %s", root, exc)
        return found

    for path in children:
        if path.is_dir():
            # Legacy: still accept D:\Attendance\NOIDA\file.xls if someone used old layout
            name = path.name.upper()
            if name in ALLOWED_LOCATIONS or name == "HYDERABAD":
                loc = ALLOWED_LOCATIONS.get(name, "HYD")
                for child in sorted(path.iterdir()):
                    if not _is_excel_report(child):
                        continue
                    resolved = child.resolve()
                    if resolved in seen:
                        continue
                    seen.add(resolved)
                    found.append((child, loc))
            continue
        if not _is_excel_report(path):
            continue
        resolved = path.resolve()
        if resolved in seen:
            continue
        seen.add(resolved)
        found.append((path, office_location))

    return found


def _discover_drop_files() -> list:
    """Scan ATTENDANCE_DIR only (LOCATION_CODE from .env)."""
    found = []
    seen = set()
    office_location = normalize_location(validate_location(get_location_from_config()))

    drop = _office_drop_dir()
    if drop is None:
        LOGGER.warning("ATTENDANCE_DIR is not set; nothing to scan")
        return found

    LOGGER.info("Scanning drop folder %s as %s", drop, office_location)
    found.extend(_scan_flat_drop(drop, office_location, seen))
    return found


def _notify_failure(summary: dict) -> None:
    if summary.get("empty") or summary.get("already_running") or summary.get("ok"):
        return
    url = os.getenv("DDO_ALERT_WEBHOOK_URL", "").strip()
    if not url:
        LOGGER.error(
            "Import failed (processed=%s failed=%s). Set DDO_ALERT_WEBHOOK_URL to notify a channel.",
            summary.get("processed"),
            summary.get("failed"),
        )
        return
    payload = json.dumps(
        {
            "text": (
                "DDO attendance import failed: "
                f"processed={summary.get('processed')} failed={summary.get('failed')} "
                f"skipped={summary.get('skipped')}"
            ),
            "processed": summary.get("processed"),
            "failed": summary.get("failed"),
            "skipped": summary.get("skipped"),
            "results": (summary.get("results") or [])[:20],
        }
    ).encode("utf-8")
    request = urllib.request.Request(
        url,
        data=payload,
        method="POST",
        headers={"Content-Type": "application/json", "Accept": "application/json"},
    )
    try:
        with urllib.request.urlopen(request, timeout=10) as response:
            LOGGER.info("Failure alert posted (%s)", response.status)
    except Exception as exc:
        LOGGER.error("Could not post failure alert: %s", exc)


def _write_last_run(summary: dict) -> Path:
    output_dir = Path(CONFIG["output_dir"])
    output_dir.mkdir(parents=True, exist_ok=True)
    payload = {
        "finished_at": datetime.now().isoformat(timespec="seconds"),
        **summary,
    }
    path = output_dir / "last_run.json"
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    _notify_failure(summary)
    return path


def process_drop_folder(explicit_location=None) -> dict:
    _configure_logging()
    Path(CONFIG["processed_dir"]).mkdir(parents=True, exist_ok=True)
    Path(CONFIG["failed_dir"]).mkdir(parents=True, exist_ok=True)

    lock = _SingleInstance(_lock_path())
    if not lock.acquire():
        summary = {
            "ok": True,
            "already_running": True,
            "processed": 0,
            "failed": 0,
            "skipped": 0,
            "empty": False,
        }
        LOGGER.info("Importer already running; exiting")
        _write_last_run(summary)
        return summary

    try:
        candidates = _discover_drop_files()
        known_hashes = _load_processed_hashes()
        processed = 0
        failed = 0
        skipped = 0
        results = []

        if not candidates:
            LOGGER.info("Drop folder is empty; nothing to import")
            summary = {
                "ok": True,
                "already_running": False,
                "processed": 0,
                "failed": 0,
                "skipped": 0,
                "empty": True,
            }
            _write_last_run(summary)
            return summary

        for path, folder_location in candidates:
            location = explicit_location or folder_location
            try:
                location = normalize_location(location)
            except ValueError:
                LOGGER.warning("Skipping %s; unknown location %s", path.name, location)
                skipped += 1
                continue
            if not _is_file_ready(path):
                skipped += 1
                continue
            try:
                file_hash = _file_hash(path)
            except OSError as exc:
                LOGGER.warning("Skipping %s; could not read file: %s", path.name, exc)
                skipped += 1
                continue
            if file_hash in known_hashes:
                LOGGER.info("Skipping %s; identical file already imported", path.name)
                move_to_processed(str(path))
                skipped += 1
                results.append({"file": path.name, "location": location, "status": "DUPLICATE"})
                continue

            batch = process_file(str(path), location=location, raise_on_error=False)
            api_ok = bool(batch and (batch.api_result or {}).get("ok"))
            if batch is None or batch.status != "SUCCESS" or not api_ok:
                failed += 1
                results.append(
                    {
                        "file": path.name,
                        "location": location,
                        "status": "FAILED",
                        "error": None if batch is None else batch.error,
                    }
                )
                continue
            _remember_processed_hash(file_hash, path.name, location)
            known_hashes[file_hash] = True
            processed += 1
            file_status = "PARTIAL" if (batch.api_result or {}).get("partial") else "SUCCESS"
            results.append({"file": path.name, "location": location, "status": file_status})

        summary = {
            "ok": failed == 0,
            "already_running": False,
            "processed": processed,
            "failed": failed,
            "skipped": skipped,
            "empty": False,
            "results": results,
        }
        _write_last_run(summary)
        LOGGER.info(
            "Drop-folder run complete: processed=%s failed=%s skipped=%s",
            processed,
            failed,
            skipped,
        )
        return summary
    finally:
        try:
            _prune_output_dir()
            _prune_folder_keep_newest(CONFIG["processed_dir"], CONFIG.get("archive_keep", 15))
            _prune_folder_keep_newest(CONFIG["failed_dir"], CONFIG.get("archive_keep", 15))
        except Exception as exc:
            LOGGER.warning("archive prune failed: %s", exc)
        lock.release()


def _parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description="Import attendance Excel reports and POST them to the DDO++ API.",
    )
    parser.add_argument(
        "files",
        nargs="*",
        help="Attendance .xls/.xlsx files. If omitted, scans ATTENDANCE_DIR.",
    )
    parser.add_argument(
        "--location",
        help="Override location for this run: AGRA, NOIDA, or HYD.",
    )
    return parser.parse_args(argv)


if __name__ == "__main__":
    _configure_logging()
    args = _parse_args()
    location = normalize_location(args.location) if args.location else None

    files = [Path(item) for item in args.files] if args.files else []
    if not files:
        default_file = _default_input_file()
        if default_file is not None:
            files = [default_file]
        else:
            summary = process_drop_folder(explicit_location=location)
            print(
                "DDO importer: processed={processed} failed={failed} skipped={skipped} empty={empty}".format(
                    **summary
                )
            )
            sys.exit(0 if summary.get("ok") else 1)

    missing = [path for path in files if not path.exists()]
    if missing:
        LOGGER.error("Input file not found: %s", missing[0])
        sys.exit(
            "Usage: python scripts/importer_script.py | "
            "python scripts/importer_script.py <attendance.xls>"
        )

    for path in files:
        file_location = location or infer_location(path) or get_location_from_config()
        process_file(str(path), location=file_location)
    try:
        _prune_output_dir()
        _prune_folder_keep_newest(CONFIG["processed_dir"], CONFIG.get("archive_keep", 15))
        _prune_folder_keep_newest(CONFIG["failed_dir"], CONFIG.get("archive_keep", 15))
    except Exception as exc:
        LOGGER.warning("archive prune failed: %s", exc)
